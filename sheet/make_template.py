"""Builds the master spreadsheet (to import into Google Sheets) from the verified casebook."""
import json, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

S = "/tmp/claude-0/-home-claude/f827d55f-71c0-5ac4-9168-2237972221a3/scratchpad"
rows = json.load(open(f"{S}/casebook_rows.json"))
geo = {}
for i in range(1, 5):
    for g in json.load(open(f"{S}/maps/geo_out_{i}.json")):
        geo[g["no"]] = g
LABEL_FIX = {'AG-177': 'Western US firewall reboots', 'AG-097': 'Brazil blackout hacking claim', 'AG-262': 'Berlin power cable arsons',
             'AG-188': 'Australia state-actor announcement', 'AG-212': 'China time-service intrusion claim',
             'AG-058': 'Police-staged MK explosions', 'AG-280': 'Leyton warehouse arson'}
notes = json.load(open(f"{S}/maps/notes.json"))

COLS = [
    ("id", 10, "Catalogue number, AG- followed by three digits. Never reuse a number."),
    ("status", 11, "published = shown on the site; draft or review = kept off the site; withdrawn = removed."),
    ("added", 12, "Date the entry went live (YYYY-MM-DD). New entries with the latest dates go into the email digest."),
    ("title", 42, "Full title of the event."),
    ("label", 28, "Short map label, three to five words."),
    ("dates", 22, "Dates as they should be printed."),
    ("year", 7, "First year of the event (a number). Places the event in its decade."),
    ("place", 26, "Country and place as printed."),
    ("lat", 9, "Latitude of the main marker, decimal degrees."),
    ("lon", 9, "Longitude of the main marker, decimal degrees."),
    ("more_sites", 20, "Further sites of the same event: 'lat, lon; lat, lon'. Leave empty if none."),
    ("precision", 11, "site, area or approximate."),
    ("case_type", 15, "missed, obvious, false-alarm, wrongly-blamed or still-open."),
    ("scope", 16, "civilian service, or adjacent for commercial or defense targets."),
    ("context", 22, "armed conflict, short of armed conflict, terrorism / non-state, criminal, no hostile act."),
    ("services", 26, "One or more from the Lists tab, separated by commas."),
    ("strength", 9, "1, 2 or 3."),
    ("summary", 60, "Two to four sentences on what happened."),
    ("records", 60, "Who held which part of the record."),
    ("gap", 60, "Why the records stayed apart."),
    ("joining", 60, "Who joined them, when, with what authority."),
    ("harm_found_later", 45, "Harm identified only by later reconstruction (accountability gap). Optional."),
    ("use_in_volume", 45, "What the event would let a chapter show."),
    ("sources", 60, "One source per line: citation | URL"),
    ("verification", 40, "verified, or partly verified: what remains open."),
    ("map_note", 45, "Optional one-sentence note shown on the decade map."),
]

wb = Workbook()
ws = wb.active
ws.title = "Events"
F = "Arial"
hdr_font = Font(name=F, bold=True, size=10, color="FFFFFF")
hdr_fill = PatternFill("solid", fgColor="5A1A24")
body_font = Font(name=F, size=10)
thin = Side(style="thin", color="D9D9D9")
for c, (name, width, help_) in enumerate(COLS, 1):
    cell = ws.cell(row=1, column=c, value=name)
    cell.font, cell.fill = hdr_font, hdr_fill
    cell.alignment = Alignment(vertical="center")
    cell.comment = Comment(help_, "Template")
    ws.column_dimensions[cell.column_letter].width = width
ws.row_dimensions[1].height = 20
ws.freeze_panes = "E2"

for i, r in enumerate(rows, 2):
    g = geo.get(r["no"], {})
    pts = g.get("points") or []
    vals = {
        "id": r["no"], "status": "published", "added": datetime.date(2026, 10, 22),
        "title": r["title"], "label": LABEL_FIX.get(r["no"], g.get("label", "")),
        "dates": r["dates"], "year": r["year"], "place": r["place"],
        "lat": pts[0][0] if pts else None, "lon": pts[0][1] if pts else None,
        "more_sites": "; ".join(f"{p[0]}, {p[1]}" for p in pts[1:]),
        "precision": g.get("precision", ""), "case_type": r["case_type"],
        "scope": "adjacent" if r.get("scope") == "adjacent" else "civilian service",
        "context": r["context"], "services": ", ".join(r["services"]), "strength": r["strength"],
        "summary": r["summary"], "records": r["records"], "gap": r["gap"], "joining": r["joining"],
        "harm_found_later": r.get("accountability_gap") or "", "use_in_volume": r.get("book_use") or "",
        "sources": "\n".join(f"{s.get('cite','').strip()} | {s.get('url','').strip()}" for s in r.get("sources", [])),
        "verification": r["verification"], "map_note": notes.get(r["no"], ""),
    }
    for c, (name, _, _) in enumerate(COLS, 1):
        cell = ws.cell(row=i, column=c, value=vals[name])
        cell.font = body_font
        cell.alignment = Alignment(vertical="top", wrap_text=name in ("title", "summary", "records", "gap", "joining", "harm_found_later", "use_in_volume", "sources", "verification", "map_note"))
        cell.border = Border(bottom=thin)
        if name == "added": cell.number_format = "yyyy-mm-dd"
    ws.row_dimensions[i].height = 60
last = len(rows) + 1
ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(COLS)).column_letter}{last}"

# Lists tab with allowed values
L = wb.create_sheet("Lists")
lists = {
    "status": ["published", "draft", "review", "withdrawn"],
    "case_type": ["missed", "obvious", "false-alarm", "wrongly-blamed", "still-open"],
    "scope": ["civilian service", "adjacent"],
    "context": ["armed conflict", "short of armed conflict", "terrorism / non-state", "criminal", "no hostile act"],
    "precision": ["site", "area", "approximate"],
    "strength": [1, 2, 3],
    "services": ["water", "energy", "food and agriculture", "health", "telecommunications and internet", "transport", "navigation and space", "other"],
}
for c, (k, v) in enumerate(lists.items(), 1):
    h = L.cell(row=1, column=c, value=k); h.font, h.fill = hdr_font, hdr_fill
    L.column_dimensions[h.column_letter].width = 30
    for j, x in enumerate(v, 2):
        L.cell(row=j, column=c, value=x).font = body_font
colidx = {name: i for i, (name, _, _) in enumerate(COLS, 1)}
for c, k in enumerate(lists, 1):
    if k == "services":
        continue  # several values per cell; checked by the site build instead
    col = L.cell(row=1, column=c).column_letter
    dv = DataValidation(type="list", formula1=f"=Lists!${col}$2:${col}${len(lists[k]) + 1}", allow_blank=True,
                        showErrorMessage=True, errorTitle="Not in the list", error=f"Choose a value for {k} from the Lists tab.")
    ws.add_data_validation(dv)
    letter = ws.cell(row=1, column=colidx[k]).column_letter
    dv.add(f"{letter}2:{letter}2000")

# Editor notes: private, never published
N = wb.create_sheet("Editor notes")
for c, (h, w) in enumerate([("id", 10), ("note", 80), ("by", 18), ("date", 12)], 1):
    cell = N.cell(row=1, column=c, value=h); cell.font, cell.fill = hdr_font, hdr_fill
    N.column_dimensions[cell.column_letter].width = w
ex = ["AG-177", "Example row: operator identity still unconfirmed; check NERC lesson document before naming.", "SS", datetime.date(2026, 10, 22)]
for c, v in enumerate(ex, 1):
    cell = N.cell(row=2, column=c, value=v); cell.font = Font(name=F, size=10, italic=True, color="7F7F7F")
N.cell(row=2, column=4).number_format = "yyyy-mm-dd"

# README
R = wb.create_sheet("README", 0)
R.column_dimensions["A"].width = 110
text = [
    ("Aggregation Gap Casebook: master spreadsheet", True),
    ("This workbook feeds the public website. The site rebuilds once a day from the Events tab.", False),
    ("", False),
    ("How to edit", True),
    ("1. Edit cells in the Events tab directly. Hover over a column header to see what it holds.", False),
    ("2. To publish a new event: give it the next free AG number, set status to published and set added to today's date.", False),
    ("3. To take an event off the site, set status to withdrawn. Do not delete rows or reuse numbers.", False),
    ("4. Put coordinates in decimal degrees (lat, lon). Further sites go in more_sites as 'lat, lon; lat, lon'.", False),
    ("5. Sources: one per line, written as 'citation | URL'.", False),
    ("6. Values for status, case_type, scope, context, precision and strength come from the Lists tab.", False),
    ("", False),
    ("What is public", True),
    ("Only the Events tab is published to the website, and every cell in it is public, drafts included.", False),
    ("Never put private notes, reviewers' comments or the emails of people who suggest events in the Events tab.", False),
    ("Private remarks go in the Editor notes tab. Suggestions from the public arrive in the form's own responses tab.", False),
    ("", False),
    ("Updates", True),
    ("The website rebuilds automatically every morning. To publish a change at once, run the 'Build and deploy' workflow on GitHub.", False),
    ("Subscribers receive an email digest of entries whose added date is newer than the last digest.", False),
]
for i, (t, bold) in enumerate(text, 1):
    c = R.cell(row=i, column=1, value=t)
    c.font = Font(name=F, size=12 if bold else 10, bold=bold)
    c.alignment = Alignment(wrap_text=True, vertical="top")

out = "/home/claude/aggregation-gap-site/sheet/Aggregation_Gap_Casebook_master.xlsx"
wb.save(out)
print("saved", out, "rows", len(rows))
